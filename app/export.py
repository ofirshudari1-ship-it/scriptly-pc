"""ייצוא הקלטה (סיכום + תמלול), או ייצוא מרוכז (digest) של כמה הקלטות יחד.
תומך בארבעה פורמטים - Word (.docx, ברירת מחדל), טקסט פשוט (.txt), Markdown (.md)
ו-SRT/כתוביות (.srt, הקלטה בודדת בלבד) - לפי `export_format` בהגדרות. כולם נכתבים
מקומית לגמרי, בלי שום תלות ברשת."""
import re
from pathlib import Path
from typing import List

from docx import Document
from docx.shared import RGBColor

from .i18n import t
from .stats import compute_stats, compute_talk_ratio, compute_tone_breakdown, compute_top_tags, format_duration
from .tasks import get_all_tasks

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
NAVY = RGBColor(0x15, 0x28, 0x47)
TEAL = RGBColor(0x00, 0x99, 0x91)

# Single-meeting export supports SRT (needs per-segment timestamps); the digest
# (multi-meeting) export does not, since a caption file only makes sense for one
# recording at a time.
EXPORT_EXTENSIONS = {"docx": ".docx", "txt": ".txt", "md": ".md", "srt": ".srt"}
DIGEST_EXPORT_EXTENSIONS = {"docx": ".docx", "txt": ".txt", "md": ".md"}


def _strip_markdown(text: str) -> str:
    """הופך Markdown פשוט לטקסט רגיל קריא - לכותרות (`#`) ולתבליטים (`-`/`*`) בלבד,
    זה כל מה שהסיכומים שלנו בפועל מייצרים (ראו SYSTEM_PROMPT ב-summarize.py)."""
    out = []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("#"):
            out.append(s.lstrip("#").strip().upper())
        elif s.startswith("- ") or s.startswith("* "):
            out.append("  • " + BOLD_RE.sub(r"\1", s[2:].strip()))
        else:
            out.append(BOLD_RE.sub(r"\1", s))
    return "\n".join(out)


def _add_markdown_paragraph(doc: Document, line: str):
    stripped = line.strip()
    if stripped.startswith("- ") or stripped.startswith("* "):
        p = doc.add_paragraph(style="List Bullet")
        stripped = stripped[2:]
    else:
        p = doc.add_paragraph()
    pos = 0
    for m in BOLD_RE.finditer(stripped):
        p.add_run(stripped[pos : m.start()])
        p.add_run(m.group(1)).bold = True
        pos = m.end()
    p.add_run(stripped[pos:])


def export_meeting_docx(meeting, out_path: Path) -> Path:
    doc = Document()

    title = doc.add_heading(meeting.title, level=1)
    title.runs[0].font.color.rgb = NAVY

    if meeting.tags:
        tag_p = doc.add_paragraph()
        tag_p.add_run(", ".join(f"#{tg}" for tg in meeting.tags)).italic = True

    if meeting.summary:
        for line in meeting.summary.splitlines():
            s = line.lstrip("#").strip()
            if line.startswith("#") and s:
                doc.add_heading(s, level=2)
            elif s:
                _add_markdown_paragraph(doc, line)

    if meeting.notes:
        doc.add_heading(t("notes_label"), level=2)
        doc.add_paragraph(meeting.notes)

    if meeting.transcript:
        doc.add_heading(t("tab_transcript"), level=2)
        for line in meeting.transcript.splitlines():
            if line.strip():
                doc.add_paragraph(line)

    doc.save(str(out_path))
    return out_path


def export_meeting_txt(meeting, out_path: Path) -> Path:
    parts = [meeting.title, "=" * len(meeting.title), ""]
    if meeting.tags:
        parts += [", ".join(f"#{tg}" for tg in meeting.tags), ""]
    if meeting.summary:
        parts += [_strip_markdown(meeting.summary), ""]
    if meeting.notes:
        parts += [t("notes_label") + ":", meeting.notes, ""]
    if meeting.transcript:
        parts += [t("tab_transcript") + ":", "-" * 40, meeting.transcript]
    out_path.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    return out_path


def export_meeting_md(meeting, out_path: Path) -> Path:
    parts = [f"# {meeting.title}", ""]
    if meeting.tags:
        parts += [", ".join(f"#{tg}" for tg in meeting.tags), ""]
    if meeting.summary:
        parts += [meeting.summary, ""]
    if meeting.notes:
        parts += [f"## {t('notes_label')}", "", meeting.notes, ""]
    if meeting.transcript:
        parts += [f"## {t('tab_transcript')}", "", "```", meeting.transcript, "```"]
    out_path.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    return out_path


def _srt_timestamp(seconds: float) -> str:
    seconds = max(0.0, seconds)
    total_ms = int(round(seconds * 1000))
    hours, total_ms = divmod(total_ms, 3_600_000)
    minutes, total_ms = divmod(total_ms, 60_000)
    secs, ms = divmod(total_ms, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{ms:03d}"


def export_meeting_srt(meeting, out_path: Path) -> Path:
    """Subtitle/caption export (.srt) - a format every competitor transcription tool
    (Otter/Fireflies/Fathom) offers so the transcript can be dropped straight into a
    video editor as captions. Built from the real per-segment timestamps captured
    during transcription (app/transcribe.py, persisted as segments.json by
    controller.py) when available; older recordings that predate segments.json fall
    back to spreading transcript lines evenly across the recorded duration."""
    segments = getattr(meeting, "segments", None) or []
    if not segments:
        lines = [ln for ln in (meeting.transcript or "").splitlines() if ln.strip()]
        if not lines or meeting.duration_seconds <= 0:
            out_path.write_text("", encoding="utf-8")
            return out_path
        per_line = meeting.duration_seconds / len(lines)
        segments = [{"start": i * per_line, "end": (i + 1) * per_line, "text": ln} for i, ln in enumerate(lines)]

    blocks = []
    index = 1
    for seg in segments:
        text = (seg.get("text") or "").strip()
        if not text:
            continue
        blocks.append(f"{index}\n{_srt_timestamp(seg['start'])} --> {_srt_timestamp(seg['end'])}\n{text}\n")
        index += 1
    out_path.write_text("\n".join(blocks), encoding="utf-8")
    return out_path


def export_meeting(meeting, out_path: Path, fmt: str = "docx") -> Path:
    """מפצל לפי הפורמט הנבחר בהגדרות (`export_format`) - עד לתיקון הזה, הדשבורד קרא
    תמיד ל-`export_meeting_docx` ישירות בלי קשר לבחירת המשתמש, כך שבחירת TXT/MD
    בהגדרות לא הייתה משפיעה בפועל על שום ייצוא."""
    if fmt == "txt":
        return export_meeting_txt(meeting, out_path)
    if fmt == "md":
        return export_meeting_md(meeting, out_path)
    if fmt == "srt":
        return export_meeting_srt(meeting, out_path)
    return export_meeting_docx(meeting, out_path)


def _summary_preview(summary: str, max_lines: int = 4) -> List[str]:
    """כמה השורות הראשונות בעלות תוכן מתוך הסיכום (בלי הכותרת "## תקציר" עצמה) - לתקציר תמציתי בדוח המרוכז."""
    lines = []
    for line in summary.splitlines():
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        lines.append(s)
        if len(lines) >= max_lines:
            break
    return lines


def export_digest_docx(meetings: List, out_path: Path, range_label: str) -> Path:
    """דוח מרוכז על כמה הקלטות יחד: תמונת מצב (סטטיסטיקה, טון, תגיות), כל המשימות הפתוחות,
    ואז תקציר קצר לכל הקלטה - במקום לפתוח כל הקלטה בנפרד."""
    doc = Document()

    title = doc.add_heading(f"{t('digest_title')} — {range_label}", level=1)
    title.runs[0].font.color.rgb = NAVY
    doc.add_paragraph(t("digest_generated_note"))

    # --- תמונת מצב ---
    doc.add_heading(t("digest_overview"), level=2)
    for line in _digest_overview_lines(meetings):
        doc.add_paragraph(line, style="List Bullet")

    # --- משימות פתוחות ---
    open_tasks = [tsk for tsk in get_all_tasks(meetings) if not tsk.done]
    doc.add_heading(t("tasks_header"), level=2)
    if open_tasks:
        for tsk in open_tasks:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{tsk.text} ")
            run = p.add_run(f"({tsk.meeting_title})")
            run.italic = True
            run.font.color.rgb = TEAL
    else:
        doc.add_paragraph(t("tasks_all_done"))

    # --- תקציר לכל הקלטה ---
    doc.add_heading(t("digest_per_meeting"), level=2)
    for m in meetings:
        doc.add_heading(m.title, level=3)
        if m.tags:
            tag_p = doc.add_paragraph()
            tag_p.add_run(", ".join(f"#{tg}" for tg in m.tags)).italic = True
        if not m.summary:
            doc.add_paragraph(t("processing_marker"))
            continue
        for line in _summary_preview(m.summary):
            _add_markdown_paragraph(doc, line)

    doc.save(str(out_path))
    return out_path


def _digest_overview_lines(meetings: List) -> List[str]:
    s = compute_stats(meetings)
    lines = [
        f"{t('stats_total_recordings')}: {s['total']}",
        f"{t('stats_total_time')}: {format_duration(s['total_seconds']) if s['total_seconds'] else t('stats_none')}",
        f"{t('stats_avg_length')}: {format_duration(s['avg_seconds']) if s['avg_seconds'] else t('stats_none')}",
    ]
    ratio = compute_talk_ratio(meetings)
    if ratio:
        lines.append(f"{t('stats_talk_ratio')}: {t('speaker_you')} {ratio['you_pct']}% · {t('speaker_other')} {ratio['other_pct']}%")
    tones = compute_tone_breakdown(meetings)
    if any(tones.values()):
        lines.append(
            f"{t('stats_tone')}: {t('tone_positive')} {tones['positive']} · {t('tone_neutral')} {tones['neutral']} · {t('tone_tense')} {tones['tense']}"
        )
    top_tags = compute_top_tags(meetings)
    if top_tags:
        lines.append(f"{t('stats_top_tags')}: " + ", ".join(f"#{tg} ({c})" for tg, c in top_tags))
    return lines


def export_digest_txt(meetings: List, out_path: Path, range_label: str) -> Path:
    title = f"{t('digest_title')} — {range_label}"
    parts = [title, "=" * len(title), "", t("digest_generated_note"), ""]
    parts += [t("digest_overview"), "-" * len(t("digest_overview"))]
    parts += [f"  • {line}" for line in _digest_overview_lines(meetings)]
    parts.append("")

    open_tasks = [tsk for tsk in get_all_tasks(meetings) if not tsk.done]
    parts += [t("tasks_header"), "-" * len(t("tasks_header"))]
    if open_tasks:
        parts += [f"  • {tsk.text} ({tsk.meeting_title})" for tsk in open_tasks]
    else:
        parts.append(t("tasks_all_done"))
    parts.append("")

    parts += [t("digest_per_meeting"), "-" * len(t("digest_per_meeting"))]
    for m in meetings:
        parts.append("")
        parts.append(m.title)
        if m.tags:
            parts.append(", ".join(f"#{tg}" for tg in m.tags))
        if not m.summary:
            parts.append(t("processing_marker"))
            continue
        parts.append(_strip_markdown("\n".join(_summary_preview(m.summary))))

    out_path.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    return out_path


def export_digest_md(meetings: List, out_path: Path, range_label: str) -> Path:
    parts = [f"# {t('digest_title')} — {range_label}", "", t("digest_generated_note"), ""]
    parts += [f"## {t('digest_overview')}", ""]
    parts += [f"- {line}" for line in _digest_overview_lines(meetings)]
    parts.append("")

    open_tasks = [tsk for tsk in get_all_tasks(meetings) if not tsk.done]
    parts += [f"## {t('tasks_header')}", ""]
    if open_tasks:
        parts += [f"- {tsk.text} _{tsk.meeting_title}_" for tsk in open_tasks]
    else:
        parts.append(t("tasks_all_done"))
    parts.append("")

    parts += [f"## {t('digest_per_meeting')}", ""]
    for m in meetings:
        parts.append(f"### {m.title}")
        if m.tags:
            parts.append(", ".join(f"#{tg}" for tg in m.tags))
        if not m.summary:
            parts.append(t("processing_marker"))
            continue
        parts.append("\n".join(_summary_preview(m.summary)))
        parts.append("")

    out_path.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    return out_path


def export_digest(meetings: List, out_path: Path, range_label: str, fmt: str = "docx") -> Path:
    if fmt == "txt":
        return export_digest_txt(meetings, out_path, range_label)
    if fmt == "md":
        return export_digest_md(meetings, out_path, range_label)
    return export_digest_docx(meetings, out_path, range_label)
